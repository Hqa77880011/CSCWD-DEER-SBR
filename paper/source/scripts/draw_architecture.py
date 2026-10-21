"""Generate editable PowerPoint Fig. 1 and export its vector PDF and PNG.

Dependencies: python-pptx, PyMuPDF, LibreOffice, Liberation Sans.
Run: python scripts/draw_architecture.py
The diagram follows main.tex Eqs. (1)-(12); it contains no raster artwork.
"""
from pathlib import Path
import json
import subprocess
import tempfile
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.dml.color import RGBColor
from pptx.oxml.xmlchemy import OxmlElement
import pymupdf as fitz

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'figures'
W, H = 18.3, 5.35
C = {'ink':'233347','muted':'526275','line':'4B5D70','frame':'BAC6D2',
     'blue':'28649B','bluefill':'EBF2F9','violet':'79539C','violetfill':'F2EDF7',
     'green':'287960','greenfill':'EAF4EE','obs':'69798A','white':'FFFFFF'}
prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
prs.core_properties.title = 'Figure 1. Evidence Routing and Supported Residuals'
prs.core_properties.subject = 'Editable architecture diagram for the anonymous manuscript'
prs.core_properties.author = 'Anonymous Authors'
prs.core_properties.last_modified_by = ''
prs.core_properties.keywords = 'DEER, SBR, cooperative perception, LiDAR, radar'
prs.core_properties.comments = 'Native editable shapes and connectors. Generated from scripts/draw_architecture.py.'
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = RGBColor.from_string(C['white'])


def rect(name, x, y, w, h, fill, border, width=1.1, rounded=False):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name = name
    if rounded: sh.adjustments[0] = .08
    sh.fill.solid(); sh.fill.fore_color.rgb = RGBColor.from_string(C.get(fill,fill))
    sh.line.color.rgb = RGBColor.from_string(C.get(border,border)); sh.line.width = Pt(width)
    return sh


def text(name,x,y,w,h,txt,size=20,bold=False,color='ink',align=PP_ALIGN.CENTER):
    sh = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    sh.name = name
    tf=sh.text_frame; tf.clear(); tf.word_wrap=False; tf.auto_size=MSO_AUTO_SIZE.NONE
    tf.margin_left=tf.margin_right=Inches(.05)
    tf.margin_top=tf.margin_bottom=0
    tf.vertical_anchor=MSO_ANCHOR.MIDDLE
    for i,line in enumerate(txt.split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.alignment=align; p.space_before=p.space_after=Pt(0)
        p.font.name='Arial'; p.font.size=Pt(size); p.font.bold=bold
        p.font.color.rgb=RGBColor.from_string(C.get(color,color))
        p.line_spacing=1.05
    return sh


def edge(name, points, color='line', dash=None, arrow=True, width=3.2):
    for i,((x1,y1),(x2,y2)) in enumerate(zip(points,points[1:])):
        sh=slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
        sh.name=f'{name}_{i}'
        sh.line.color.rgb=RGBColor.from_string(C[color]); sh.line.width=Pt(width)
        if dash: sh.line.dash_style=dash
        if arrow and i==len(points)-2:
            a=OxmlElement('a:tailEnd'); a.set('type','triangle'); a.set('w','lg'); a.set('len','lg')
            sh._element.spPr.get_or_add_ln().append(a)


rect('local_group',.12,.64,13.12,4.55,'FFFFFF','frame')
rect('cooperative_group',13.55,.64,4.6,4.55,'FFFFFF','frame')
text('local_heading',.25,.05,12.85,.45,'LOCAL PROCESSING AT EACH AGENT',23,True,'ink')
text('cooperative_heading',13.6,.05,4.45,.45,'COOPERATIVE DETECTION',21,True,'green')
# Feature nodes preserve the forward graph and original radar bypass.
y,h=2.0,1.5
nodes=[('encoders',.4,2.3,'LiDAR / Radar','Pillar encoders\n+ scatter','bluefill','blue'),
       ('deer',3.0,2.5,'DEER (A)','Evidence routing\nRetained gains','bluefill','blue'),
       ('base',5.9,1.6,'Base','Projection','bluefill','blue'),
       ('sbr',7.85,2.9,'SBR (B)','Local donors + null\nSupport + bound','violetfill','violet'),
       ('output',11.1,1.65,'Output','Projection','bluefill','blue'),
       ('attfuse',13.8,2.1,'Backbone','+ AttFuse','greenfill','green'),
       ('boxes',16.4,1.45,'Vehicle','3D boxes','greenfill','green')]
for name,x,w,title,body,fill,border in nodes:
    rect(name,x,y,w,h,fill,border,1.5,True)
    text(name+'_title',x+.07,y+.13,w-.14,.48,title,19.5 if name == 'encoders' else 22,True,border)
    text(name+'_body',x+.07,y+.66,w-.14,.69,body,18.5)
for (name,x,w,*_), (n2,x2,w2,*_) in zip(nodes,nodes[1:]):
    edge(name+'_to_'+n2,[(x+w,y+h/2),(x2,y+h/2)])
text('base_feature',7.48,2.13,.39,.44,'Z',19,False,'ink')
text('corrected_feature',10.72,2.13,.39,.44,'Z′',19,False,'ink')
text('agent_features',12.78,2.19,.93,.35,'{Xᵢ}',19,False,'ink')
# Detached type control enters SBR independently of the routed feature path.
edge('detached_types',[(4.2,2),(4.2,1.4),(9.23,1.4),(9.23,2)],'violet',MSO_LINE_DASH_STYLE.DASH)
text('type_label',4.29,.84,4.95,.49,'Detached missing / spurious scores',19,False,'violet')
# Shared radar feature bus supplies both the donor path and final concatenation.
edge('radar_bypass',[(1.55,3.5),(1.55,4.02),(11.93,4.02),(11.93,3.5)],'blue')
edge('radar_to_sbr',[(9.65,4.02),(9.65,3.5)],'blue')
text('radar_label',4.73,4.06,3.40,.43,'Original radar feature',18,False,'blue')
# Current occupancy/FOV/availability maps have a separate dotted control bus.
rect('observation_maps',.4,4.56,3.85,.43,'FFFFFF','FFFFFF',0)
text('observation_label',.34,4.52,3.95,.51,'Current observation maps Ω',18,False,'obs')
edge('observations_bus',[(4.3,4.78),(8.28,4.78),(8.28,3.5)],'obs',MSO_LINE_DASH_STYLE.ROUND_DOT,True,2.2)
edge('observations_to_deer',[(4.6,4.78),(4.6,3.5)],'obs',MSO_LINE_DASH_STYLE.ROUND_DOT,True,2.2)
# White bridges distinguish crossings from joins on the feature bus.
for x in [4.6,8.28]:
    rect(f'crossing_gap_{x}',x-.055,3.95,.11,.14,'FFFFFF','FFFFFF',0)
    edge(f'observation_bridge_{x}',[(x,3.95),(x,4.09)],'obs',MSO_LINE_DASH_STYLE.ROUND_DOT,False,2.2)
# Legends explain line semantics without relying on color alone.
legend=[('Features','line',None,3.96),('Detached scores','violet',MSO_LINE_DASH_STYLE.DASH,4.37),
        ('Observation maps','obs',MSO_LINE_DASH_STYLE.ROUND_DOT,4.78)]
for label,col,dash,ly in legend:
    edge('legend_'+label,[(13.88,ly),(14.6,ly)],col,dash,True,2.6)
    text('legend_text_'+label,14.72,ly-.22,3.25,.44,label,18,False,'muted',PP_ALIGN.LEFT)
text('training_note',.40,.94,3.5,.47,'Type targets: training only',17.5,False,'muted',PP_ALIGN.LEFT)
# Notes remain editable and travel with the PowerPoint source.
notes=slide.notes_slide.notes_text_frame
notes.text=('Fig. 1: Local LiDAR/radar pillar features pass through DEER routing, base projection, '
            'SBR correction and output projection before shared AttFuse aggregation and detection. '
            'The original radar feature supplies SBR and the shared output bypass. Detached missing/spurious '
            'scores reach SBR separately. Current observation maps provide occupancy, FOV, range and availability. '
            'Source: manuscript main.tex, Method. All scientific relationships are unchanged from the implementation. '
            'Author identity: Anonymous Authors. Designed for full-width placement in the six-page IEEE manuscript.')
OUT.mkdir(exist_ok=True)
# Suppress theme effects: default Office shadows otherwise rasterize in PDF.
for sh in slide.shapes:
    sh._element.spPr.append(OxmlElement('a:effectLst'))
    for effect in sh._element.xpath('.//a:effectRef'):
        effect.set('idx','0')
# Check text against fixed boxes with an additional 10% width reserve.
# PDF Base-14 Helvetica is metrically compatible with Arial for these labels.
for sh in slide.shapes:
    if not sh.has_text_frame or not sh.text.strip():
        continue
    tf=sh.text_frame
    inner_width=(sh.width-tf.margin_left-tf.margin_right)/12700
    inner_height=(sh.height-tf.margin_top-tf.margin_bottom)/12700
    total_height=0
    for paragraph in tf.paragraphs:
        size=paragraph.font.size.pt if paragraph.font.size else 20
        font=fitz.Font('hebo' if paragraph.font.bold else 'helv')
        width=font.text_length(paragraph.text,fontsize=size)
        assert width <= inner_width*.90, (sh.name,'width',width,inner_width)
        total_height += size*1.18
    assert total_height <= inner_height, (sh.name,'height',total_height,inner_height)
pptx=OUT/'architecture.pptx'; prs.save(pptx)
with tempfile.TemporaryDirectory(prefix='cscwd_fig_export_') as temp:
    temp=Path(temp)
    result = subprocess.run(['libreoffice','-env:UserInstallation='+ (temp/'profile').as_uri(), '--headless',
                    '--convert-to','pdf','--outdir',str(temp),str(pptx)],check=True,capture_output=True,text=True)
    if not (temp/'architecture.pdf').exists():
        raise RuntimeError('LibreOffice did not export the figure: ' + result.stdout + result.stderr)
    doc=fitz.open(temp/'architecture.pdf')
    doc.set_metadata({'title':prs.core_properties.title,'author':'Anonymous Authors',
                      'subject':'DEER/SBR architecture; native PowerPoint vector export',
                      'keywords':prs.core_properties.keywords,'creator':'PowerPoint figure generator','producer':'LibreOffice'})
    doc.save(OUT/'architecture.pdf',garbage=4,deflate=True)
    doc[0].get_pixmap(matrix=fitz.Matrix(2,2),alpha=False).save(OUT/'architecture.png')
    (OUT/'architecture.svg').write_text(doc[0].get_svg_image(text_as_path=False))
    assert len(doc)==1
    assert not doc[0].get_images(), 'Figure must remain vector-only'
    assert 'SBR (B)' in doc[0].get_text()
print('Created editable PPTX, vector PDF/SVG, and PNG preview:',pptx)
